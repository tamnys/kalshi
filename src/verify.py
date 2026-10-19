"""Numerical/accounting checks tied to the research's acceptance criteria."""
import hashlib
import json
from decimal import Decimal
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar
from audit import ROOT,DEFAULT_SOURCE,name_score
from kelly import fit_firth,kelly_fraction,path_metrics


def run(source_path=DEFAULT_SOURCE):
    source=json.loads((ROOT/'results/source_inventory.json').read_text())
    assert hashlib.sha256(source_path.read_bytes()).hexdigest()==source['sha256']
    original=load_workbook(source_path,data_only=True).active
    cleaned=load_workbook(ROOT/'data/cleaned/MMA betting - audited.xlsx',data_only=True)
    assert original['I16'].value=='y' and cleaned['Bets']['I16'].value=='cancelled'
    assert original['I63'].value=='y' and cleaned['Bets']['I63'].value=='n'
    assert cleaned['Bets']['A82'].value.year==2025 and cleaned['Bets']['A96'].value.year==2025
    assert cleaned['Bets']['F18'].value is None and cleaned['Bets']['N18'].value=='leal -900'
    assert name_score('you','Youssef Zalal')<1 and name_score('yoo','JooSang Yoo')==1
    d=pd.read_csv(ROOT/'data/cleaned/kalshi_fights.csv')
    q=d[d.kalshi_price.notna()&d.payout.notna()]
    independently_calculated=[]
    for r in q.itertuples():
        price=Decimal(str(r.kalshi_price));payout=Decimal(str(r.payout))
        fee=Decimal('0.07')*price*(1-price)
        independently_calculated.append(float(payout/(price+fee)-1))
    assert np.allclose(independently_calculated,q.kalshi_return)
    observed=pd.read_csv(ROOT/'artifacts/private/kalshi_observed_prices.csv')
    assert (observed.candle_end<=observed.card_start).all()
    assert not observed.ticker.duplicated().any()
    joined=observed.merge(d[['row','payout']],on='row')
    assert np.allclose(joined.settlement,joined.payout)
    # Analytic binary Kelly versus direct maximization of expected log wealth.
    for p,c in [(0.6,0.5),(0.4,0.35),(0.75,0.7),(0.4,0.5)]:
        f=float(kelly_fraction(p,c))
        objective=lambda v:-(p*np.log1p(v*(1/c-1))+(1-p)*np.log1p(-v))
        optimum=minimize_scalar(objective,bounds=(0,1),method='bounded')
        assert objective(f)<=optimum.fun+np.finfo(float).eps*len(q)
    metric=path_metrics(np.array([.1,.1]),np.array([1.,-1.]))
    assert np.isclose(metric['final_wealth'],.99) and np.isclose(metric['max_drawdown'],.1)
    beta=fit_firth(np.ones((4,1)),np.array([0.,0.,1.,1.]))
    assert np.allclose(beta,0)
    assert np.isfinite(fit_firth(np.ones((4,1)),np.ones(4))).all()
    pred=pd.read_csv(ROOT/'artifacts/private/prequential_predictions.csv')
    assert pred.groupby('date',sort=False).ngroups==15
    summary=json.loads((ROOT/'results/analysis_summary.json').read_text())
    assert summary['binary_fights']==185 and summary['wins']==104
    assert np.isclose(summary['profit'],d[d.y.notna()&d.kalshi_cost.notna()].kalshi_return.sum())
    # Quotes, rather than already-decided prices, support every observed row.
    assert observed.ask.between(0,1,inclusive='neither').all()
    print('PASS: source preservation, reviewed corrections, identity matching, quote timing, settlement sides, fee arithmetic, Kelly optimum, compounding, finite separated fits, and aggregate reconciliation.')


if __name__=='__main__':run()
