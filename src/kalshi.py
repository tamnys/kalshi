"""Recover public historical Kalshi asks, then estimate missing asks without outcomes.

Use the last valid ask candle ending at/before ESPN's card start. Never use a
settled market's last_price, which leaks the result. Archive and live endpoints
are both searched, and every response is cached for reproduction.
"""
import hashlib
import json
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen

import numpy as np
import pandas as pd
from scipy.special import expit, logit
from openpyxl import load_workbook
from audit import ROOT, name_score, norm

BASE = 'https://api.elections.kalshi.com/trade-api/v2'
CACHE = ROOT / 'data/raw/kalshi'


def get(path, params=None):
    url = BASE + path + ('?' + urlencode(params) if params else '')
    CACHE.mkdir(exist_ok=True, parents=True)
    file = CACHE / (hashlib.sha256(url.encode()).hexdigest() + '.json')
    if not file.exists():
        data = json.load(urlopen(url))
        file.write_text(json.dumps(dict(url=url, response=data)))
    return json.loads(file.read_text())['response'], url


def markets():
    allmarkets = []
    for path in ('/markets', '/historical/markets'):
        # API-documented maximum page size, not an analysis cap.
        params = dict(series_ticker='KXUFCFIGHT', limit=1000)
        if path == '/markets':
            params['status'] = 'settled'
        while True:
            data, _ = get(path, params)
            allmarkets.extend(dict(m, archived=path.startswith('/historical')) for m in data['markets'])
            print(path, 'markets fetched', len(allmarkets), flush=True)
            if not data.get('cursor'):
                break
            params['cursor'] = data['cursor']
    return allmarkets


def fit_mapping(x, y, model):
    lx, ly = logit(x), logit(y)
    if model == 'identity':
        return np.array([0., 1.])
    if model == 'shift':
        return np.array([np.mean(ly - lx), 1.])
    return np.linalg.lstsq(np.column_stack([np.ones(len(x)), lx]), ly, rcond=None)[0]


def run():
    df = pd.read_csv(ROOT / 'data/cleaned/fights.csv')
    mm = markets()
    series, _ = get('/series/KXUFCFIGHT')
    rows = []
    for r in df.itertuples():
        if r.status == 'cancelled':
            continue
        datecode = pd.Timestamp(r.date).strftime('%y%b%d').upper()
        candidates = [m for m in mm if ('-' + datecode) in m['ticker']]
        # All candidates on the correct date are checked; exact normalized name
        # containment or reviewed transliterations establish the selected side.
        pname = 'you' if r.row == 33 else {'sumadaerji': 'sumudaerji'}.get(r.pick, r.pick)
        scores = sorted([(name_score(pname, m.get('yes_sub_title', '')), m) for m in candidates],
                        key=lambda x: x[0], reverse=True)
        if not scores or scores[0][0] < 1:
            # Do not auto-match fuzzy prices; print for explicit review if needed.
            if scores:
                print('UNMATCHED', r.row, r.pick, [(s, m['yes_sub_title']) for s, m in scores[:2]], flush=True)
            continue
        tied = [m for s, m in scores if s == 1]
        if len(tied) != 1:
            print('AMBIGUOUS', r.row, r.pick, [m['ticker'] for m in tied], flush=True)
            continue
        m = tied[0]
        sb = json.loads((ROOT / f'data/raw/scoreboards/{r.date}.json').read_text())['events'][0]
        end = int(pd.Timestamp(sb['date']).timestamp())
        start = int(pd.Timestamp(m['open_time']).timestamp())
        if start >= end:
            print('OPENED_AFTER_CARD_START', r.row, flush=True)
            continue
        path = (f'/historical/markets/{m["ticker"]}/candlesticks' if m['archived'] else
                f'/series/KXUFCFIGHT/markets/{m["ticker"]}/candlesticks')
        # Hour candles cover the full available pre-card history, with no arbitrary
        # staleness rejection. Candle age is retained for sensitivity analyses.
        data, url = get(path, dict(start_ts=start, end_ts=end, period_interval=60))
        valid = []
        for c in data['candlesticks']:
            ask = c.get('yes_ask', {}).get('close_dollars', c.get('yes_ask', {}).get('close'))
            bid = c.get('yes_bid', {}).get('close_dollars', c.get('yes_bid', {}).get('close'))
            if ask is None:
                continue
            ask, bid = float(ask), float(bid) if bid is not None else np.nan
            # Legacy integer-cent fields versus documented dollar-string fields.
            if isinstance(c['yes_ask'].get('close'), int) and 'close_dollars' not in c['yes_ask']:
                ask /= 100
                bid /= 100
            if 0 < ask < 1 and c['end_period_ts'] <= end:
                valid.append((c, ask, bid))
        if not valid:
            print('NO_PRECARD_ASK', r.row, m['ticker'], flush=True)
            continue
        c, ask, bid = max(valid, key=lambda t: t[0]['end_period_ts'])
        rows.append(dict(row=r.row, date=r.date, ticker=m['ticker'], fighter=m['yes_sub_title'],
                         ask=ask, bid=bid, candle_end=c['end_period_ts'], card_start=end,
                         age_minutes=(end-c['end_period_ts'])/60,
                         settlement=float(m.get('settlement_value_dollars') or (m['result']=='yes')),
                         source=url, market_open=m['open_time']))
        print('PRICE', r.row, m['ticker'], ask, bid, 'age', (end-c['end_period_ts'])/60, flush=True)
    prices = pd.DataFrame(rows)
    prices.to_csv(ROOT / 'artifacts/private/kalshi_observed_prices.csv', index=False)
    df = df.merge(prices[['row', 'ask', 'bid', 'ticker', 'age_minutes']], on='row', how='left')
    overlap = df[df.ask.notna() & df.venue.eq('sportsbook_unspecified') & df.price.notna()].copy()
    # Do not train against derived American odds in the Kalshi quote rows.
    x, y = overlap.price.to_numpy(), overlap.ask.to_numpy()
    errors = []
    for model in ('identity', 'shift', 'affine_logit'):
        preds = np.empty(len(overlap))
        for date in overlap.date.unique():
            test = overlap.date.to_numpy() == date
            coef = fit_mapping(x[~test], y[~test], model)
            preds[test] = expit(coef[0] + coef[1]*logit(x[test]))
        errors.append(dict(model=model, rmse=float(np.sqrt(np.mean((preds-y)**2))),
                           mae=float(np.mean(np.abs(preds-y))),
                           bias=float(np.mean(preds-y)),
                           coefficients=fit_mapping(x,y,model).tolist()))
        overlap['cv_' + model] = preds
    winner = min(errors, key=lambda e: e['rmse'])
    a, b = winner['coefficients']
    # Restore the two opposite-side rows by estimating the *quoted opponent's*
    # bid, then complementing it. Bid mapping is learned from independent pairs.
    # This is an imputation, not recovery of the missing selected-fighter odds.
    bid_overlap = overlap[overlap.bid.between(0,1,inclusive='neither')]
    abid, bbid = fit_mapping(bid_overlap.price.to_numpy(), bid_overlap.bid.to_numpy(), winner['model'])
    df['kalshi_price'] = np.nan
    df['price_source'] = ''
    for i, r in df.iterrows():
        if r.status == 'cancelled':
            df.loc[i, 'price_source'] = 'cancelled'
        elif pd.notna(r.ask):
            df.loc[i, ['kalshi_price', 'price_source']] = [r.ask, 'api_precard_ask']
        elif pd.notna(r.kalshi_text):
            df.loc[i, ['kalshi_price', 'price_source']] = [r.price, 'workbook_kalshi_quote']
        elif r.status == 'missing_pick_price':
            amer = r.american
            opponent_q = 1 / (1 + (amer/100 if amer > 0 else 100/abs(amer)))
            estimated = 1-expit(abid+bbid*logit(opponent_q))
            df.loc[i, ['kalshi_price', 'price_source']] = [estimated, 'imputed_from_opponent']
        else:
            df.loc[i, ['kalshi_price', 'price_source']] = [expit(a+b*logit(r.price)), 'imputed_from_sportsbook']
    df['kalshi_cost'] = df.kalshi_price + .07*df.kalshi_price*(1-df.kalshi_price)
    # A No purchase against the selected fighter uses 1 minus its bid, not
    # 1 minus its ask. Missing bid quotes use the separately fitted bid mapping.
    df['kalshi_no_price'] = np.nan
    for i, r in df.iterrows():
        if r.status == 'cancelled':
            continue
        if pd.notna(r.bid) and 0 < r.bid < 1:
            no_price = 1-r.bid
        elif r.status == 'missing_pick_price':
            amer=r.american
            oq=1/(1+(amer/100 if amer>0 else 100/abs(amer)))
            no_price=expit(a+b*logit(oq))
        else:
            no_price=1-expit(abid+bbid*logit(r.price))
        df.loc[i,'kalshi_no_price']=no_price
    df['kalshi_no_cost']=df.kalshi_no_price+.07*df.kalshi_no_price*(1-df.kalshi_no_price)
    df['kalshi_no_return']=(1-df.payout)/df.kalshi_no_cost-1
    df['kalshi_return'] = df.payout/df.kalshi_cost-1
    df['kalshi_gross_return'] = df.payout/df.kalshi_price-1
    df.to_csv(ROOT / 'data/cleaned/kalshi_fights.csv', index=False)
    overlap.to_csv(ROOT / 'artifacts/private/price_mapping_overlap.csv', index=False)
    residual = overlap.ask-overlap['cv_' + winner['model']]
    summary = dict(indexed_market_count=len(mm), price_source_counts=df.price_source.value_counts().to_dict(),
                   oldest_indexed_market=min(m['open_time'] for m in mm),
                   overlap_n=len(overlap), overlap_cards=overlap.date.nunique(),
                   models=errors, chosen_model=winner['model'],
                   residual_quantiles={str(q):float(residual.quantile(q)) for q in (.025,.25,.5,.75,.975)},
                   opponent_bid_coefficients=[float(abid),float(bbid)],
                   largest_quote_age_minutes=float(prices.age_minutes.max()),
                   series_fee_multiplier=series['series']['fee_multiplier'],
                   notes=['All prices are pre-card asks, not last traded settled values.',
                          'Price-model fitting never uses fight outcomes.',
                          'Leave-one-card-out RMSE chooses among identity, shift and affine-logit mapping.',
                          'Imputed prices before market inception are counterfactual, not historical fills.',
                          'Price reconstruction learned from later cards is not a strictly prospective backtest.'])
    (ROOT/'results/kalshi_price_reconstruction.json').write_text(json.dumps(summary,indent=2)+'\n')
    wb=load_workbook(ROOT/'data/cleaned/MMA betting - audited.xlsx')
    tab=wb.create_sheet('Kalshi analysis data')
    cols=['row','date','event','fight','pick','confidence','result','price_source','kalshi_price','kalshi_cost','kalshi_return','kalshi_no_price','kalshi_no_cost','kalshi_no_return','ticker','age_minutes']
    tab.append(cols)
    for vals in df[cols].itertuples(index=False,name=None):
        tab.append([None if pd.isna(v) else v for v in vals])
    wb.save(ROOT/'data/cleaned/MMA betting - audited.xlsx')
    print(json.dumps(summary,indent=2),flush=True)


if __name__ == '__main__':
    run()
