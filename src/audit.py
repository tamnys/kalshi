"""Inventory the source workbook and independently reconcile recorded outcomes.

Downloads ESPN scoreboards to a local cache. Fuzzy matches are diagnostic only:
the cleaning step uses explicitly reviewed corrections, never fuzzy rewrites.
"""
import argparse
import hashlib
import json
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from urllib.request import Request, urlopen

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path('/Users/j/Code/Kalshi/MMA betting - UFC Edmonton Oct 17 2026 - results updated.xlsx')
EXPECTED_SHA256 = '32f5219d41c182159dcf53de4e019d703d7ad04b061b3e2505444dd68aa427b3'


def norm(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)


def read_source(path):
    if hashlib.sha256(path.read_bytes()).hexdigest() != EXPECTED_SHA256:
        raise ValueError('This analysis has reviewed cell-specific corrections for one immutable workbook. Re-audit a different source before applying them.')
    ws = load_workbook(path, data_only=True).active
    rows = []
    date = event = None
    for row in ws.iter_rows(min_row=2):
        v = [c.value for c in row]
        if not v[2]:
            continue
        if v[0]:
            date = pd.Timestamp(v[0])
            # Two demonstrable year typos, independently verified against UFC.
            if row[0].row in (82, 96):
                date = date.replace(year=2025)
            event = v[1] or 'UFC Fight Night: Emmett vs Murphy'
        rows.append(dict(row=row[0].row, date=date.date().isoformat(), event=event,
                         fight=v[2], pick=str(v[3]).strip(), confidence=v[4],
                         odds_text=v[5], kalshi_text=v[6], conversion=v[7], result=v[8],
                         market_url=v[9], displayed_quote=v[10],
                         verified_result=v[11], result_url=v[12]))
    return pd.DataFrame(rows)


def cached_scoreboard(date):
    cache = ROOT / 'data/raw/scoreboards' / (date + '.json')
    url = 'https://site.api.espn.com/apis/site/v2/sports/mma/ufc/scoreboard?dates=' + date.replace('-', '')
    if not cache.exists():
        cache.parent.mkdir(parents=True, exist_ok=True)
        data = urlopen(Request(url, headers={'User-Agent': 'MMA-data-audit/1.0'})).read()
        json.loads(data)
        cache.write_bytes(data)
    return json.loads(cache.read_text()), url


def name_score(short, full):
    a, b = norm(short), norm(full)
    tokens = full.split()
    phrases = {norm(' '.join(tokens[i:j])) for i in range(len(tokens)) for j in range(i+1,len(tokens)+1)}
    if a in phrases:
        return 1.0
    return max(SequenceMatcher(None, a, norm(' '.join(tokens[i:]))).ratio() for i in range(len(tokens)))


def audit(path):
    df = read_source(path)
    out = ROOT / 'artifacts/private'
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / 'source_rows.csv', index=False)
    matches = []
    for date, group in df.groupby('date', sort=True):
        data, url = cached_scoreboard(date)
        comps = [(e['name'], k, c) for e in data.get('events', []) for k, c in enumerate(e['competitions'])]
        print(date, len(group), 'recorded;', len(comps), 'ESPN bouts')
        for r in group.itertuples():
            parts = re.split(r'\s+vs\.?\s+', r.fight.strip(), flags=re.I)
            candidates = []
            for ename, order, c in comps:
                cc = c['competitors']
                if len(cc) != 2 or len(parts) != 2:
                    continue
                names = [x['athlete']['displayName'] for x in cc]
                score = max(name_score(parts[0], names[0]) + name_score(parts[1], names[1]),
                            name_score(parts[0], names[1]) + name_score(parts[1], names[0])) / 2
                pickidx = max(range(2), key=lambda i: name_score(r.pick, names[i]))
                winner = next((x['athlete']['displayName'] for x in cc if x.get('winner')), None)
                expected = ('y' if cc[pickidx].get('winner') else 'n') if winner else 'draw_or_unknown'
                candidates.append(dict(row=r.row, fight=r.fight, pick=r.pick, recorded=r.result,
                                       match_score=score, official_fighters=' / '.join(names),
                                       winner=winner, expected=expected, espn_order=order,
                                       competition_id=c['id'], source=url, event=ename,
                                       completed=c.get('status', {}).get('type', {}).get('completed')))
            if candidates:
                best = max(candidates, key=lambda x: x['match_score'])
                matches.append(best)
                if best['match_score'] != 1 or best['recorded'] != best['expected']:
                    print(' REVIEW', json.dumps(best))
            else:
                matches.append(dict(row=r.row, fight=r.fight, expected='unavailable', source=url))
    pd.DataFrame(matches).to_csv(out / 'outcome_reconciliation.csv', index=False)
    meta = dict(filename=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                fight_rows=len(df), events=df.date.nunique(), results=df.result.value_counts().to_dict(),
                confidences={str(k): int(v) for k, v in df.confidence.value_counts().sort_index().items()},
                first_date=df.date.min(), last_date=df.date.max(),
                kalshi_quote_rows=int(df.kalshi_text.notna().sum()),
                independent_outcome_source='ESPN scoreboards; reviewed corrections also cite UFC')
    (ROOT / 'results').mkdir(exist_ok=True)
    (ROOT / 'results/source_inventory.json').write_text(json.dumps(meta, indent=2) + '\n')
    print(json.dumps(meta, indent=2))


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    audit(ap.parse_args().source)
