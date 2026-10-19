"""Apply reviewed, source-backed corrections; preserve original cells in a log."""
import argparse
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from audit import DEFAULT_SOURCE, ROOT, read_source

SOURCES = {
    'cancelled': 'https://www.ufc.com/news/updates-ufc-313-pereira-vs-ankalaev',
    'result': 'https://www.ufc.com/news/main-card-results-highlights-winner-interviews-ufc-fight-night-moreno-vs-erceg-mexico-city',
    'ufc314': 'https://www.ufc.com/news/official-judges-scorecards-ufc-314-volkanovski-vs-lopes-kaseya-center-miami',
    'kansas': 'https://www.ufc.com/event/ufc-fight-night-april-26-2025',
    'emmett': 'https://www.ufc.com/event/ufc-fight-night-april-05-2025',
    'draw': 'https://api.elections.kalshi.com/trade-api/v2/series/KXUFCFIGHT',
    'fees': 'https://kalshi.com/docs/kalshi-fee-schedule.pdf',
}


def clean(source):
    rawdir, cleandir = ROOT / 'data/raw', ROOT / 'data/cleaned'
    rawdir.mkdir(exist_ok=True, parents=True)
    cleandir.mkdir(exist_ok=True, parents=True)
    shutil.copy2(source, rawdir / source.name)
    df = read_source(source)
    wb = load_workbook(source)
    ws = wb.active
    log = []

    def change(cell, value, reason, evidence):
        old = ws[cell].value
        if old == value:
            return
        log.append(dict(cell=cell, old=str(old), new=str(value), reason=reason, source=evidence))
        ws[cell] = value

    change('I16', 'cancelled', 'Canceled at UFC 313; not a settled win. April rescheduling remains a separate prediction.', SOURCES['cancelled'])
    change('I63', 'n', 'Kevin Borjas defeated Ronaldo Rodriguez.', SOURCES['result'])
    change('A82', datetime(2025, 4, 12), 'Correct event year.', SOURCES['ufc314'])
    change('A96', datetime(2025, 4, 26), 'Correct event year.', SOURCES['kansas'])
    change('B69', 'UFC Fight Night: Emmett vs Murphy', 'Fill missing event heading from date and fight card.', SOURCES['emmett'])
    # Cosmetic changes are verified against the independently reconciled fight card.
    spellings = {3: ('temurov', 'temirov'), 4: ('aldritch', 'aldrich'),
                 33: ('yoo', 'you'), 69: ('demopolous', 'demopoulos'),
                 73: ('frumza', 'frunza'), 76: ('osborne', 'osbourne'),
                 82: ('cornole', 'cornolle'), 84: ('sumadaerji', 'sumudaerji'),
                 157: ('cerquiera', 'cerqueira'), 170: ('shabazyan', 'shahbazyan'),
                 181: ('zheleznyakova', 'zhelezniakova')}
    for row, (old, new) in spellings.items():
        date = df.loc[df.row.eq(row), 'date'].iloc[0].replace('-', '')
        url = 'https://site.api.espn.com/apis/site/v2/sports/mma/ufc/scoreboard?dates=' + date
        for col in ('C', 'D', 'F'):
            cell = f'{col}{row}'
            if isinstance(ws[cell].value, str) and old in ws[cell].value:
                change(cell, ws[cell].value.replace(old, new), 'Standardize verified fighter spelling.', url)
    for r in df.itertuples():
        for col in ('C', 'D', 'F'):
            cell = f'{col}{r.row}'
            if isinstance(ws[cell].value, str):
                change(cell, ws[cell].value.strip(), 'Trim whitespace.', 'Workbook cell')
    # Opponent prices cannot be sign-flipped: the other side includes an unknown spread.
    for row in (18, 40):
        ws[f'N{row}'] = ws[f'F{row}'].value
        change(f'F{row}', None, 'Quoted price names opponent; selected-fighter price unknown. Original saved in N.', 'Workbook D/F comparison')
        change(f'H{row}', None, 'Remove invalid selected-fighter conversion.', 'Workbook D/F comparison')
    ws['N1'] = 'Original invalid odds'
    ws['O1'] = 'Analysis status'
    ws['P1'] = 'Price used / breakeven'
    ws['Q1'] = 'Gross profit per dollar at risk'
    ws['R1'] = 'Exact net win odds'
    df.loc[df.row.eq(16), 'result'] = 'cancelled'
    df.loc[df.row.eq(63), 'result'] = 'n'
    rec = pd.read_csv(ROOT / 'artifacts/private/outcome_reconciliation.csv').set_index('row')
    for i, r in df.iterrows():
        m = re.search(r'([+-])\s*(\d+(?:\.\d+)?)\s*$', r.odds_text)
        assert m is not None, r.row
        american = float(m[2]) * (1 if m[1] == '+' else -1)
        b = american / 100 if american > 0 else 100 / abs(american)
        price = 1 / (1 + b)
        venue = 'sportsbook_unspecified'
        if pd.notna(r.kalshi_text):
            price = float(re.search(r'(\d+)\s*[¢%]', r.kalshi_text)[1]) / 100
            b = 1 / price - 1
            venue = 'kalshi'
        status = 'eligible'
        if r.row in (18, 40):
            b = price = np.nan
            status = 'missing_pick_price'
        elif r.result == 'cancelled':
            status = 'cancelled'
        elif r.result == 'draw':
            status = 'draw_separate'
        y = 1.0 if r.result == 'y' else 0.0 if r.result == 'n' else np.nan
        payout = y if r.result in ('y', 'n') else .5 if r.result == 'draw' else np.nan
        gross_return = payout / price - 1 if pd.notna(price) and pd.notna(payout) else np.nan
        # Costs are per dollar of total cash at risk, fee included. Fee rounding is
        # negligible only in the scalable-continuous-stake approximation used here.
        fee = .07 * price * (1 - price) if venue == 'kalshi' else 0.0
        actual_cost = price + fee
        net_return = payout / actual_cost - 1 if pd.notna(price) and pd.notna(payout) else np.nan
        df.loc[i, ['american', 'b', 'price', 'y', 'payout', 'gross_return', 'net_return']] = [american, b, price, y, payout, gross_return, net_return]
        df.loc[i, ['venue', 'status']] = [venue, status]
        df.loc[i, 'order'] = rec.loc[r.row, 'espn_order'] if r.row != 16 else np.nan
        df.loc[i, 'official_fighters'] = rec.loc[r.row, 'official_fighters'] if r.row != 16 else None
        ws[f'O{r.row}'] = status
        if pd.notna(price):
            ws[f'P{r.row}'] = price
            ws[f'R{r.row}'] = b
            # Preserve the supplied rounded H display; exact analytics are in R.
            if pd.notna(gross_return):
                ws[f'Q{r.row}'] = gross_return
    df['confidence'] = df.confidence.astype(int)
    df = df.sort_values(['date', 'order', 'row']).reset_index(drop=True)
    df['sequence'] = np.arange(len(df))
    assert not df.loc[df.status.ne('cancelled')].duplicated(['date', 'official_fighters']).any()
    assert df.loc[df.status.eq('eligible'), ['price', 'confidence', 'y']].notna().all().all()
    assert df.loc[df.status.eq('eligible'), 'price'].between(0, 1, inclusive='neither').all()
    # All completed rows were reviewed, including all fuzzy spelling matches.
    for r in df.itertuples():
        if r.status == 'cancelled':
            continue
        expected = rec.loc[r.row, 'expected']
        if r.result in ('y', 'n'):
            assert r.result == expected, r.row
    df.to_csv(cleandir / 'fights.csv', index=False)
    pd.DataFrame(log).to_csv(cleandir / 'corrections.csv', index=False)
    sheet = wb.create_sheet('Corrections')
    sheet.append(['Cell', 'Original', 'Corrected', 'Reason', 'Evidence'])
    for entry in log:
        sheet.append(list(entry.values()))
    info = wb.create_sheet('Analysis notes')
    for line in [
        'Source preserved unchanged; corrections logged in Corrections tab.',
        'R is the exact net win payout per dollar staked; H is the supplied rounded display.',
        'P uses exact Kalshi cents when provided; otherwise derives breakeven from American odds.',
        'Q excludes fees; draw at row 180 pays $0.50 per contract under KXUFCFIGHT rules.',
        'The two opponent-price rows are excluded from priced inference; do not flip their odds.',
        'Dates/event labels inherit within a card in the normalized CSV only.',
        'Fight order follows ESPN scoreboard order; worksheet ordering is not always chronological.',
        'Confidence 1–5 is ordinal, not a stated probability.',
        'Fees, fills, pre-fight timestamps and full-record completeness are not proven by the workbook.',
    ]:
        info.append([line])
    for sheet in (ws, wb['Corrections'], info):
        sheet.freeze_panes = 'A2'
        for c in sheet[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E78')
    wb.save(cleandir / 'MMA betting - audited.xlsx')
    counts = df.status.value_counts().to_dict()
    summary = dict(status_counts=counts, corrected_win_to_loss=1, cancelled_marked_as_win=1,
                   year_corrections=2, missing_event_filled=1, missing_pick_prices=2,
                   logged_cell_changes=len(log), sources=SOURCES,
                   binary_priced_fights=int(df.status.eq('eligible').sum()),
                   original_rounded_conversion_max_absolute_error=float(np.max(np.abs(
                       df.loc[df.status.eq('eligible') & df.venue.eq('sportsbook_unspecified'), 'b'] -
                       df.loc[df.status.eq('eligible') & df.venue.eq('sportsbook_unspecified'), 'conversion']))))
    (ROOT / 'results/cleaning_summary.json').write_text(json.dumps(summary, indent=2) + '\n')
    print(json.dumps(summary, indent=2))
    d = df[df.status.eq('eligible')]
    for name, g in [('all', d), ('favorites', d[d.price > .5]), ('underdogs', d[d.price < .5])]:
        print(name, len(g), int(g.y.sum()), g.gross_return.sum(), g.gross_return.mean())
        print(g.groupby('confidence').agg(n=('y','size'), wins=('y','sum'), profit=('gross_return','sum'), roi=('gross_return','mean')))


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    clean(ap.parse_args().source)
